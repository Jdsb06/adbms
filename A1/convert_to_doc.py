import docx
import os
import re

def clean_xml_string(s):
    # Filter out control characters that violate XML specification
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', s)

def markdown_to_docx(md_path, docx_path):
    doc = docx.Document()
    with open(md_path, 'r', encoding='utf-8', errors='ignore') as f:
        lines = f.readlines()

    base_dir = os.path.dirname(os.path.abspath(md_path))

    for line in lines:
        stripped = clean_xml_string(line.strip())
        if not stripped:
            continue
        
        # Check for image syntax ![alt](filename.png)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            img_filename = img_match.group(2)
            img_path = os.path.join(base_dir, img_filename)
            if os.path.exists(img_path):
                doc.add_paragraph(clean_xml_string(img_match.group(1)))
                doc.add_picture(img_path, width=docx.shared.Inches(6.0))
            continue

        # Check headings
        if stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith('- '):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        else:
            clean_text = re.sub(r'[*_`$]', '', stripped)
            doc.add_paragraph(clean_text)

    doc.save(docx_path)
    print(f"Successfully converted {md_path} -> {docx_path}")

if __name__ == '__main__':
    markdown_to_docx('Assignment_Report.md', 'Assignment_Report.docx')
    markdown_to_docx('LLM_Conversations.md', 'LLM_Conversations.docx')

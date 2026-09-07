import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def md_to_docx(md_file_path, docx_file_path):
    print(f"Converting {md_file_path} to {docx_file_path}...")
    doc = Document()

    # Set page margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    with open(md_file_path, 'r', encoding='utf-8', errors='ignore') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []
    in_code = False
    code_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle Code blocks
        if stripped.startswith('```'):
            if in_code:
                code_text = ''.join(code_lines)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(code_text)
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                in_code = False
                code_lines = []
            else:
                in_code = True
                code_lines = []
            continue

        if in_code:
            code_lines.append(line)
            continue

        # Handle Markdown Tables
        if stripped.startswith('|') and '|' in stripped:
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(stripped)
            continue
        else:
            if in_table:
                if table_lines:
                    rows = []
                    for tline in table_lines:
                        cols = [c.strip() for c in tline.split('|')[1:-1]]
                        rows.append(cols)

                    if len(rows) > 1 and all(set(c) <= set(':- ') for c in rows[1]):
                        header = rows[0]
                        body = rows[2:]
                    else:
                        header = rows[0]
                        body = rows[1:]

                    num_cols = len(header)
                    table = doc.add_table(rows=len(body) + 1, cols=num_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER

                    # Format Header Row
                    hdr_cells = table.rows[0].cells
                    for i, h_text in enumerate(header):
                        cell_text = re.sub(r'[*`\$\\]', '', h_text)
                        hdr_cells[i].text = cell_text
                        set_cell_background(hdr_cells[i], "1F497D")
                        p = hdr_cells[i].paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.font.size = Pt(9.5)

                    # Format Body Rows
                    for r_idx, row_data in enumerate(body):
                        row_cells = table.rows[r_idx + 1].cells
                        bg_color = "F2F2F2" if r_idx % 2 == 1 else "FFFFFF"
                        for c_idx in range(num_cols):
                            val = row_data[c_idx] if c_idx < len(row_data) else ""
                            cell_text = re.sub(r'[*`\$\\]', '', val)
                            row_cells[c_idx].text = cell_text
                            set_cell_background(row_cells[c_idx], bg_color)
                            p = row_cells[c_idx].paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx in [1, 2] else WD_ALIGN_PARAGRAPH.CENTER
                            for run in p.runs:
                                run.font.size = Pt(9)
                                run.font.name = 'Calibri'

                    doc.add_paragraph().paragraph_format.space_after = Pt(6)
                    table_lines = []
                    in_table = False

        # Handle Images
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            alt_text = img_match.group(1)
            img_rel_path = img_match.group(2)
            img_path = os.path.normpath(os.path.join(os.path.dirname(md_file_path), img_rel_path))
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(5.8))
                
                caption_p = doc.add_paragraph()
                caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                c_run = caption_p.add_run(f"Figure: {alt_text}")
                c_run.font.italic = True
                c_run.font.size = Pt(9.5)
                c_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            continue

        # Handle Headings
        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            run = h.add_run(re.sub(r'[*`\$\\]', '', stripped[2:]))
            run.font.name = 'Arial'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            continue
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            run = h.add_run(re.sub(r'[*`\$\\]', '', stripped[3:]))
            run.font.name = 'Arial'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            continue
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            run = h.add_run(re.sub(r'[*`\$\\]', '', stripped[4:]))
            run.font.name = 'Arial'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue
        elif stripped.startswith('#### '):
            h = doc.add_heading(level=4)
            run = h.add_run(re.sub(r'[*`\$\\]', '', stripped[5:]))
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            run.font.bold = True
            continue

        if not stripped:
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        clean_text = re.sub(r'[*`\$\\]', '', stripped)
        run = p.add_run(clean_text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)

    doc.save(docx_file_path)
    print(f"Document saved successfully as {docx_file_path}!")

if __name__ == '__main__':
    md_to_docx('reports/Assignment_Report.md', 'reports/Assignment_Report.docx')
    md_to_docx('reports/LLM_Conversations.md', 'reports/LLM_Conversations.docx')
